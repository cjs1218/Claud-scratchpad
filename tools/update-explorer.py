#!/usr/bin/env python3
"""Rebuild the Test Directory Explorer data (test-directory-explorer/data/) for every R code in the National Genomic
Test Directory for rare and inherited disease.

Sources: the directory spreadsheet (tests, targets, methods, commissioning), the eligibility criteria PDF (verbatim
criteria, overlapping indications, pathway, requesting specialties, family structure and scope of each test) and the
PanelApp API (green genes on each panel named in the Target/Genes column).

Usage: python3 tools/update-explorer.py [--xlsx PATH-OR-URL] [--pdf PATH-OR-URL] [--no-panels]
Needs python3 with openpyxl, pdftotext (poppler) and network access for downloads and PanelApp."""
import argparse, concurrent.futures, datetime, json, os, re, subprocess, sys, tempfile, urllib.request
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from extract_criteria import lines, heading, extract, FOOTER

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
OUT = os.path.join(ROOT, 'test-directory-explorer', 'data')
BASE = 'https://www.england.nhs.uk/wp-content/uploads/2018/08/'
XLSX = BASE + 'rare-and-inherited-disease-national-genomics-test-directory-v9.xlsx'
PDF = BASE + 'rare-and-inherited-disease-eligibility-criteria-v9.1.pdf'
UA = {'User-Agent': 'Mozilla/5.0 (test-directory-explorer update)'}
SECTIONS = ['Testing Criteria', 'Overlapping indications', 'Where in Pathway', 'Requesting Specialties',
            'Specialist Service Group', 'Associated Tests']
# Words of a family-structure cell. Stray words at either end (spill-over from the Name column) are trimmed; a cell
# with other words in the middle was misread and is left blank.
FAMILY_WORDS = re.compile(r'(?i)\(?(singleton|trio|duo|parents|only|other|or|accepted|when|necessary|multiple|affected|individuals|3|generations)\)?')


def fetch(src, dest):
    if re.match(r'^https?://', src):
        with urllib.request.urlopen(urllib.request.Request(src, headers=UA), timeout=120) as r: open(dest, 'wb').write(r.read())
    else:
        open(dest, 'wb').write(open(src, 'rb').read())
    return dest


def clean(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''


def tidy_method(m):
    """Merge spelling variants ('Small Panel', 'small panel ', 'WES or Medium Panel') into one form."""
    m = clean(m)
    return re.sub(r'\b(Small|Medium|Large) Panel\b', lambda x: x.group(1) + ' panel', m[:1].upper() + m[1:]) \
        .replace('Highly specialised', 'Highly Specialised')


def read_sheet(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    note = [clean(r[0]) for r in wb.worksheets[0].iter_rows(values_only=True) if r and r[0]]
    m = re.search(r'(\d{1,2})(?:st|nd|rd|th)? (\w+) (20\d\d) v([\d.]+)', note[0])
    version = f'v{m.group(4).rstrip("0").rstrip(".") if m.group(4).endswith(".0") else m.group(4)} ({m.group(1)} {m.group(2)} {m.group(3)})'
    ws = wb['R&ID indications']
    head = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True))]
    col = {h: i for i, h in enumerate(head) if h}
    want = ['Clinical indication ID', 'Test ID', 'Clinical Indication', 'Target/Genes', 'Test Method',
            'Commissioning category', 'Specialist test group', 'Part in the eligibility criteria document']
    missing = [w for w in want if w not in col]
    if missing: sys.exit(f'Spreadsheet columns missing: {missing}')
    change = next((h for h in head if h.startswith('Changes since')), None)
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[col['Clinical indication ID']]: continue
        g = lambda k: clean(r[col[k]])
        rows.append({'ci': g('Clinical indication ID'), 'id': g('Test ID'), 'name': g('Clinical Indication'),
                     'target': g('Target/Genes'), 'method': tidy_method(g('Test Method')),
                     'commissioning': tidy_method(g('Commissioning category')), 'group': g('Specialist test group'),
                     'part': g('Part in the eligibility criteria document'),
                     'change': clean(r[col[change]]) if change else ''})
    return version, change, rows


def pdf_version(L):
    for _, l in L[:400]:
        m = re.search(r'(\d{1,2})(?:st|nd|rd|th)? ([A-Z][a-z]+) (20\d\d) v([\d.]+)\s*$', l.strip())
        if m: return f'v{m.group(4)} ({m.group(1)} {m.group(2)} {m.group(3)})'
    return ''


def block(L, start, codes_at):
    """Lines of one indication: from its heading to the next indication, Part heading or the closing change log."""
    end = min([i for i in codes_at if i > start] + [len(L)])
    for k in range(start + 1, end):
        if re.match(r'^(Part [IVXL]+\.|Change Log for)', L[k][1].strip()): end = k; break
    return [(pn, l) for pn, l in L[start:end] if l.strip() and not re.fullmatch(r'\d{1,3}', l.strip()) and not FOOTER.search(l.strip())]


def bullets(seg, short=False):
    """Bulleted lines -> (items, notes). Indented lines continue the item above; other lines are notes.
    short=True (specialty lists): every line is its own item, as some bullets lost their glyph in the PDF, unless it
    continues an item ending in a comma; footnotes ('*...') and 'For R184.1 ...' sub-headings are notes."""
    items, notes, ind = [], [], None
    for _, l in seg:
        s = l.strip()
        if s.startswith('Referrals for testing will') or s.startswith('Please note'): break
        b = re.match(r'^[•\-−·]\s*(.*)$', s)
        here = len(l) - len(l.lstrip())
        if short:
            t = b.group(1) if b else s
            if t.startswith('*') or re.match(r'For R\d', t): notes.append(t)
            elif items and not b and (items[-1].endswith(',') or t[:1].islower()): items[-1] += ' ' + t
            else: items.append(t)
            continue
        if b: items.append(b.group(1)); ind = here
        elif items and ind is not None and here > ind and not notes: items[-1] += ' ' + s
        elif not items and not notes and ind is None and not s.startswith('*'): items.append(s)
        else: notes.append(s)
    tidy = lambda xs: [re.sub(r'\s+', ' ', x).strip() for x in xs if x.strip()]
    return tidy(items), tidy(notes)


SECTION = re.compile(r'^(%s)(?: for ([\w -]+?))?(?: \(see next page\)| [–-] next page)?$' % '|'.join(SECTIONS))


def sections(seg):
    """{section: lines}; a variant such as 'Where in Pathway for Semi-Rapid Testing' is keyed 'Where in Pathway|Semi-Rapid Testing'."""
    found, cur = {}, None
    for pn, l in seg:
        m = SECTION.match(l.strip())
        if m: cur = m.group(1) + ('|' + m.group(2) if m.group(2) else ''); found[cur] = []; continue
        if cur: found[cur].append((pn, l))
    return found


def table(seg):
    """Associated Tests table -> {test id: {family, scope}} using the header's column positions."""
    out, cols, cur = {}, None, None
    for _, l in seg:
        if re.match(r'^\s*Code\s+Name\s+Optimal', l):
            names = [r'Code', r'Name', r'Optimal', r'Scope\(s\)', r'Target(?! Name)', r'Target Name', r'Method']
            cols = [re.search(n, l).start() for n in names]; cur = None; continue
        if not cols or l.strip() in ('Structure', 'Family', 'Family Structure'): continue
        m = re.match(r'^\s*(R\d+\.\d+)\b', l)
        if m: cur = out.setdefault(m.group(1), {k: '' for k in ('family', 'scope')})
        if not cur: continue
        for word in re.finditer(r'\S+', l):  # each word belongs to the column it starts in
            k = max([i for i, c in enumerate(cols) if c <= word.start() + 2] or [0])
            key = {2: 'family', 3: 'scope'}.get(k)
            if key: cur[key] = (cur[key] + ' ' + word.group(0)).strip()
    for v in out.values():
        w = re.sub(r'\s*\(\s*', ' (', v['family']).split()
        while w and not FAMILY_WORDS.fullmatch(w[0]): w.pop(0)
        while w and not FAMILY_WORDS.fullmatch(w[-1]): w.pop()
        ok = w and all(FAMILY_WORDS.fullmatch(x) for x in w) and w[0].lower() not in ('or', 'only')
        v['family'] = ' '.join(w) if ok else ''
        v['scope'] = re.sub(r'\s*,\s*', ', ', v['scope']).strip(' ,;')
        if re.search(r'\b(loci|genes?|[Pp]anel|interval|of)\b', v['scope']): v['scope'] = ''  # Target Type words: misread
    return out


def panelapp(ids):
    def one(pid):
        req = urllib.request.Request(f'https://panelapp.genomicsengland.co.uk/api/v1/panels/{pid}/',
                                     headers={'User-Agent': 'test-directory-explorer (update-explorer)', 'Accept': 'application/json'})
        for attempt in range(4):
            try:
                with urllib.request.urlopen(req, timeout=120) as r: d = json.load(r)
                break
            except Exception as e:
                if attempt == 3: print('PanelApp failed', pid, e, file=sys.stderr); return pid, None
        green = lambda key, name: sorted({name(x) for x in d.get(key, []) if str(x.get('confidence_level')) == '3'})
        e = {'name': d['name'], 'version': str(d['version']),
             'genes': green('genes', lambda x: x['gene_data']['gene_symbol'])}
        s, r = green('strs', lambda x: x['entity_name']), green('regions', lambda x: x['entity_name'])
        if s: e['strs'] = s
        if r: e['regions'] = r
        return pid, e
    with concurrent.futures.ThreadPoolExecutor(6) as ex:
        return {pid: e for pid, e in ex.map(one, sorted(ids, key=int)) if e}


def write_js(path, comment, name, obj):
    with open(path, 'w') as f:
        f.write(f'/* {comment} */\nconst {name} = ' + json.dumps(obj, ensure_ascii=False, separators=(',', ':'))
                .replace('},{"code"', '},\n{"code"').replace(',"', ',"') + ';\n')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=XLSX); ap.add_argument('--pdf', default=PDF)
    ap.add_argument('--no-panels', action='store_true', help='keep the existing panels.js')
    a = ap.parse_args()
    tmp = tempfile.mkdtemp()
    xlsx_version, change_col, rows = read_sheet(fetch(a.xlsx, os.path.join(tmp, 'td.xlsx')))
    pdf = fetch(a.pdf, os.path.join(tmp, 'td.pdf'))
    subprocess.run(['pdftotext', '-layout', pdf, os.path.join(tmp, 'td.txt')], check=True)
    L = lines(open(os.path.join(tmp, 'td.txt')).read())

    order = []
    for r in rows:
        if r['ci'] not in order: order.append(r['ci'])
    crit = extract(L, order)
    starts = {c: heading(L, c)[0] for c in order if heading(L, c)}
    at = sorted(starts.values())
    today = datetime.date.today().strftime('%-d %B %Y')
    inds, panel_ids, problems = [], set(), []
    for c in order:
        tests = [r for r in rows if r['ci'] == c]
        e = {'code': c, 'name': tests[0]['name'], 'part': tests[0]['part'], 'group': tests[0]['group'],
             'change': '; '.join(dict.fromkeys(t['change'] for t in tests if t['change'] and t['change'] != 'No change'))}
        if c in starts:
            sec = sections(block(L, starts[c], at))
            e['page'] = crit[c]['page']
            e['criteria'] = crit[c]['items']
            e['overlapping'] = [x for x in sum(bullets(sec.get('Overlapping indications', [])), [])]
            e['pathway'] = ' '.join(l.strip() for _, l in sec.get('Where in Pathway', []))
            e['specialties'], notes = bullets(sec.get('Requesting Specialties', []), True)
            if notes: e['specialtyNotes'] = notes
            e['serviceGroup'] = bullets(sec.get('Specialist Service Group', []), True)[0]
            for k in sec:
                if '|' in k and not k.startswith('Testing Criteria'):  # e.g. pathway and specialties for semi-rapid testing
                    name, variant = k.split('|')
                    lines_ = sec[k]
                    e.setdefault('variants', {}).setdefault(variant, {})[name] = (
                        ' '.join(l.strip() for _, l in lines_) if name == 'Where in Pathway' else bullets(lines_, True)[0])
            tbl = table(sec.get('Associated Tests', []))
        else:
            problems.append(f'{c}: no criteria section in the PDF'); tbl = {}
        e['tests'] = []
        for t in tests:
            panels = [int(p) for p in re.findall(r'\((\d+)\)', t['target'])]
            panel_ids |= set(panels)
            x = {k: t[k] for k in ('id', 'target', 'method', 'commissioning', 'group')}
            x.update({'panels': panels, **tbl.get(t['id'], {})})
            if t['change'] and t['change'] != 'No change': x['change'] = t['change']
            if c in starts and t['id'] not in tbl: problems.append(f'{t["id"]}: not in the PDF table')
            e['tests'].append(x)
        inds.append(e)
    inds.sort(key=lambda e: int(e['code'][1:]))
    meta = {'directory': xlsx_version, 'criteria': pdf_version(L), 'xlsxUrl': a.xlsx if a.xlsx.startswith('http') else XLSX,
            'pdfUrl': a.pdf if a.pdf.startswith('http') else PDF, 'changeColumn': change_col, 'built': today}
    os.makedirs(OUT, exist_ok=True)
    write_js(os.path.join(OUT, 'directory.js'),
             f'National Genomic Test Directory for rare and inherited disease {meta["directory"]} with eligibility criteria '
             f'{meta["criteria"]}, quoted verbatim. Generated by tools/update-explorer.py on {today}; do not edit by hand.',
             'TDX', {'meta': meta, 'indications': inds})
    if not a.no_panels:
        panels = panelapp(panel_ids)
        write_js(os.path.join(OUT, 'panels.js'),
                 f'Green (diagnostic-grade) genes, STRs and regions on each PanelApp panel named in the directory, from the '
                 f'PanelApp API on {today}. Generated by tools/update-explorer.py; do not edit by hand.',
                 'TDX_PANELS', {'date': today, 'panels': panels})
        missing = sorted(panel_ids - set(map(int, panels)))
        if missing: problems.append(f'PanelApp panels not fetched: {missing}')
    print(f'{len(inds)} indications, {sum(len(e["tests"]) for e in inds)} tests, {len(panel_ids)} panels')
    for p in problems: print('  check:', p)


if __name__ == '__main__':
    main()
